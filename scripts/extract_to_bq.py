import io
import os
import re
import time
import unicodedata
from datetime import datetime, timezone

import pandas as pd
import requests
from google.cloud import bigquery
from openpyxl import load_workbook

# O site da SGF foi refeito em Setembro de 2026. O Excel deixou de ter um URL
# fixo: passou a ser publicado com a data no nome (Historico-de-Cotacoes_MMDD.xlsx)
# dentro da pasta do mes do upload. Por isso o URL e descoberto em cada execucao
# a partir da pagina publica, em vez de estar escrito no codigo.
PAGE_URL = "https://goldensgf.pt/informacao-dos-fundos/"
MEDIA_API = "https://goldensgf.pt/wp-json/wp/v2/media"
FALLBACK_URL = "https://goldensgf.pt/wp-content/uploads/2026/09/Historico-de-Cotacoes_0916.xlsx"

# Grafia exacta como aparece na coluna "Nome do Fundo" do Excel. A comparacao
# ignora acentos e maiusculas, mas e este o valor que fica gravado no BigQuery.
FUNDS = [
    "SGF DR FINANÇAS",
    "Golden SGF Poupança Dinamica",
    "Golden SGF ETF Start",
    "Golden SGF ETF Plus",
    "PPR SGF Stoik",
    "Golden SGF TOP GESTORES",
]

# A tabela original continua a receber so o DR Financas, para nao partir nada
# que ja dependa dela. Os seis fundos vao para a tabela nova.
LEGACY_FUND = "SGF DR FINANÇAS"
TABLE_LEGACY = "sgf_dr_financas_nav"
TABLE_ALL = "sgf_navs"

MIN_ROWS_POR_FUNDO = 100

PROJECT_ID = os.environ["GCP_PROJECT_ID"]
DATASET = "PPR_SGF_DF"

HEADERS = {"User-Agent": "Mozilla/5.0"}

SCHEMA = [
    bigquery.SchemaField("data", "DATE"),
    bigquery.SchemaField("nav", "FLOAT64"),
    bigquery.SchemaField("fundo", "STRING"),
    bigquery.SchemaField("data_extracao", "TIMESTAMP"),
]


def table_ref(table):
    return f"{PROJECT_ID}.{DATASET}.{table}"


def norm(value):
    """Maiusculas, sem acentos e sem espacos a mais, para comparar nomes."""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(text.split()).upper()


def resolve_excel_url():
    """Descobre o URL actual do Excel. Tenta a pagina, depois a API do WordPress."""
    candidates = []

    try:
        print(f"A procurar o link do Excel em: {PAGE_URL}")
        html = requests.get(PAGE_URL, headers=HEADERS, timeout=60).text
        candidates += re.findall(
            r"https://goldensgf\.pt/wp-content/uploads/[^\"'\s>]+\.xlsx", html
        )
    except Exception as e:
        print(f"Nao foi possivel ler a pagina: {e}")

    if not candidates:
        try:
            print("A tentar a API de media do WordPress...")
            resp = requests.get(
                MEDIA_API,
                params={"search": "Historico", "per_page": 50,
                        "orderby": "date", "order": "desc"},
                headers=HEADERS,
                timeout=60,
            )
            resp.raise_for_status()
            candidates += [
                item.get("source_url", "")
                for item in resp.json()
                if str(item.get("source_url", "")).lower().endswith(".xlsx")
            ]
        except Exception as e:
            print(f"API de media falhou: {e}")

    # Ficar so com os ficheiros de cotacoes, mantendo a ordem de descoberta.
    cotacoes = [u for u in candidates if "COTAC" in norm(u)]
    ordered = cotacoes or candidates
    seen, urls = set(), []
    for u in ordered:
        if u not in seen:
            seen.add(u)
            urls.append(u)
    urls.append(FALLBACK_URL)

    print(f"Candidatos encontrados: {urls}")
    return urls


def download_excel(urls, retries=3):
    last_error = None
    for url in urls:
        for attempt in range(retries):
            try:
                print(f"A descarregar: {url} (tentativa {attempt + 1}/{retries})")
                response = requests.get(url, headers=HEADERS, timeout=120)
                response.raise_for_status()
                if not response.content.startswith(b"PK"):
                    raise ValueError("A resposta nao e um ficheiro xlsx.")
                print(f"Descarregado: {len(response.content)} bytes")
                return url, response.content
            except Exception as e:
                last_error = e
                print(f"Erro: {e}")
                if attempt < retries - 1:
                    time.sleep(10)
    raise RuntimeError(f"Nenhum URL funcionou. Ultimo erro: {last_error}")


def read_sheet(content):
    """Le a folha em modo streaming.

    O ficheiro novo declara ~1M de linhas, quase todas vazias. Em modo read_only
    a memoria mantem-se estavel e paramos assim que acabam os dados.
    """
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)

    def tres(row):
        """Garante exactamente 3 celulas: o read_only corta as vazias do fim."""
        cells = list(row)[:3]
        return cells + [None] * (3 - len(cells))

    header = tres(next(rows))
    print(f"Cabecalho: {header}")

    records, empty_streak = [], 0
    for row in rows:
        cells = tres(row)
        if all(c is None for c in cells):
            empty_streak += 1
            if empty_streak >= 500:
                break
            continue
        empty_streak = 0
        records.append(cells)
    wb.close()

    df = pd.DataFrame(records, columns=[str(h) for h in header])
    print(f"Linhas com dados: {len(df)}")
    return df


def pick_columns(df):
    """Identifica as colunas pelo nome; se falhar, usa a posicao."""
    by_name = {norm(c): c for c in df.columns}

    def find(*keys):
        for key in keys:
            for name, original in by_name.items():
                if key in name:
                    return original
        return None

    col_fundo = find("NOME DO FUNDO", "FUNDO") or df.columns[0]
    col_nav = find("COTACAO", "NAV", "VALOR") or df.columns[1]
    col_data = find("DATA") or df.columns[2]
    print(f"Colunas: fundo={col_fundo}, nav={col_nav}, data={col_data}")
    return col_fundo, col_nav, col_data


def transform(df):
    df.columns = [str(c).strip() for c in df.columns]
    col_fundo, col_nav, col_data = pick_columns(df)

    # norm(nome do ficheiro) -> grafia canonica que fica no BigQuery
    canonico = {norm(f): f for f in FUNDS}

    fundos_norm = df[col_fundo].map(norm)
    df_filtered = df[fundos_norm.isin(set(canonico))].copy()
    print(f"Registos dos {len(FUNDS)} fundos pedidos: {len(df_filtered)}")

    if df_filtered.empty:
        disponiveis = sorted(set(df[col_fundo].dropna().astype(str)))
        print(f"Fundos disponiveis no ficheiro: {disponiveis}")
        raise ValueError("Nenhum dos fundos pedidos foi encontrado.")

    serie_data = df_filtered[col_data]
    if pd.api.types.is_datetime64_any_dtype(serie_data):
        datas = serie_data
    else:
        datas = pd.to_datetime(serie_data, dayfirst=True, errors="coerce")

    df_out = pd.DataFrame()
    df_out["data"] = datas.dt.date
    df_out["nav"] = pd.to_numeric(df_filtered[col_nav], errors="coerce")
    df_out["fundo"] = fundos_norm.loc[df_filtered.index].map(canonico)
    df_out["data_extracao"] = datetime.now(timezone.utc)

    before = len(df_out)
    df_out = df_out.dropna(subset=["nav", "data"])
    print(f"Removidas {before - len(df_out)} linhas nulas. Total: {len(df_out)}")

    before = len(df_out)
    df_out = df_out.drop_duplicates(subset=["fundo", "data"], keep="last")
    if before != len(df_out):
        print(f"Removidas {before - len(df_out)} duplicadas (mesmo fundo e data).")

    df_out = df_out.sort_values(["fundo", "data"], ascending=[True, False])
    df_out = df_out.reset_index(drop=True)

    verificar_fundos(df_out)
    return df_out


def verificar_fundos(df):
    """Um fundo que desaparece do ficheiro tem de dar erro, nao passar em silencio."""
    resumo = df.groupby("fundo").agg(
        linhas=("data", "size"), inicio=("data", "min"), fim=("data", "max")
    )
    print("\nPor fundo:")
    print(resumo.to_string())
    print()

    em_falta = [f for f in FUNDS if f not in resumo.index]
    if em_falta:
        raise ValueError(f"Fundos nao encontrados no ficheiro: {em_falta}")

    magros = resumo[resumo["linhas"] < MIN_ROWS_POR_FUNDO]
    if not magros.empty:
        raise ValueError(
            f"Fundos com menos de {MIN_ROWS_POR_FUNDO} registos: "
            f"{list(magros.index)}. Carga abortada."
        )


def ensure_dataset_and_table(client, table):
    dataset_ref = bigquery.Dataset(f"{PROJECT_ID}.{DATASET}")
    dataset_ref.location = "EU"
    try:
        client.get_dataset(dataset_ref)
    except Exception:
        client.create_dataset(dataset_ref)
        print(f"Dataset '{DATASET}' criado.")

    ref = table_ref(table)
    try:
        client.get_table(ref)
        print(f"Tabela '{table}' ja existe.")
    except Exception:
        client.create_table(bigquery.Table(ref, schema=SCHEMA))
        print(f"Tabela '{table}' criada.")


def sanity_check(client, table, df):
    """A carga apaga e reescreve a tabela. Nao o fazer com um ficheiro suspeito."""
    ref = table_ref(table)
    try:
        existentes = client.get_table(ref).num_rows
    except Exception:
        return
    if existentes and len(df) < existentes * 0.9:
        raise ValueError(
            f"{table}: o ficheiro traz {len(df)} registos mas a tabela ja tem "
            f"{existentes}. Carga abortada para nao perder historico."
        )


def load_to_bq(client, table, df):
    ref = table_ref(table)
    job_config = bigquery.LoadJobConfig(
        schema=SCHEMA,
        write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
    )
    print(f"A carregar {len(df)} registos para {ref}...")
    client.load_table_from_dataframe(df, ref, job_config=job_config).result()
    print(f"Carga concluida. {table} tem {client.get_table(ref).num_rows} linhas.")


def main():
    client = bigquery.Client(project=PROJECT_ID)
    url, content = download_excel(resolve_excel_url())
    print(f"Fonte utilizada: {url}")

    df_todos = transform(read_sheet(content))
    df_legacy = df_todos[df_todos["fundo"] == LEGACY_FUND].reset_index(drop=True)

    for table, df in ((TABLE_ALL, df_todos), (TABLE_LEGACY, df_legacy)):
        ensure_dataset_and_table(client, table)
        sanity_check(client, table, df)
        load_to_bq(client, table, df)


if __name__ == "__main__":
    main()
